#!/usr/bin/env python3
"""md → docx/html 統一変換スクリプト(候補13)。

業務OSの文書ライフサイクル(docs/02 M-1: md生成 → 人間修正 → docx正本 → 配布)の
「md生成」直後の変換工程を1本にまとめたもの。文書タイプ(gijiroku/yoken/kaito)ごとに
体裁(見出し番号・配布先欄強調・横向き)を切り替える。

使い方:
    python3 md2office.py 入力.md --type gijiroku|yoken|kaito [--out docx,html] [--outdir DIR]

例:
    python3 md2office.py templates/議事録テンプレ.md --type gijiroku --out html
    python3 md2office.py 0710_要件定義_v1.md --type yoken --out docx,html --outdir 出力/

対応するMarkdown記法(標準ライブラリのみの簡易レンダラ):
    見出し(#〜####)/ 表(|...|)/ 箇条書き(- ,*) / 番号リスト(1. 2. ...) /
    引用(>) / 太字(**text**) / インラインコード(`code`) / 水平線(---)

依存:
    - html出力: 標準ライブラリのみで完結
    - docx出力: python-docx が入っていれば生成。無ければ案内を出してhtmlのみで継続する
      (pip install python-docx / 導入不可でもエラー終了しない)

安全原則:
    - 読み取り専用。入力.mdは一切変更しない
    - 既存の出力ファイル(docx/html)は上書きしない。既存ファイルは "_2" 等を付けて退避してから
      新しい出力を書く(手動修正済みの正本を誤って潰さないため。docs/02 M-2参照)
    - 変換して終わりにしない。実行後に必ず
      「⚠️ 正本昇格前に目視確認(docs/02 M-1)」を表示する
    - 本スクリプトはDrive・配布・送信・削除を一切行わない(ローカルファイル生成のみ)
    - レポートツールとして、想定内のエラー(入力なし/不正--type等)でも
      分かるメッセージを出して終了コード0で終わる(name_lint.pyと同じ方針)
"""

from __future__ import annotations

import argparse
import html as html_lib
import os
import re
import sys
from pathlib import Path

VALID_TYPES = ("gijiroku", "yoken", "kaito")
VALID_OUTS = ("docx", "html")

# ---------------------------------------------------------------------------
# 1. Markdown → 中間表現(ブロックのリスト)
# ---------------------------------------------------------------------------
# 各ブロックは dict。type ごとの形:
#   heading  : {type, level(1-4), runs}
#   para     : {type, runs}
#   hr       : {type}
#   quote    : {type, runs}
#   list     : {type, ordered(bool), items(list[list[runs]])}
#   table    : {type, header(list[list[runs]]), rows(list[list[list[runs]]])}

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")
HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
HR_RE = re.compile(r"^(-{3,}|\*{3,}|_{3,})\s*$")
UL_RE = re.compile(r"^[-*]\s+(.*)$")
OL_RE = re.compile(r"^\d+[.)]\s+(.*)$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def parse_inline(text: str):
    """太字・インラインコードを runs = [{"text","bold","code"}] に分解する。"""
    runs = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            runs.append({"text": part[2:-2], "bold": True, "code": False})
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            runs.append({"text": part[1:-1], "bold": False, "code": True})
        else:
            runs.append({"text": part, "bold": False, "code": False})
    return runs


def split_table_row(line: str):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def parse_markdown(text: str):
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if stripped == "":
            i += 1
            continue

        m = HEADING_RE.match(stripped)
        if m:
            level = len(m.group(1))
            blocks.append({"type": "heading", "level": level, "runs": parse_inline(m.group(2).strip())})
            i += 1
            continue

        if HR_RE.match(stripped):
            blocks.append({"type": "hr"})
            i += 1
            continue

        # コードフェンス(```): 閉じフェンスまでを整形済みテキストとして保持(インライン解釈しない)
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < n:
                i += 1  # 閉じフェンスを消費(無い場合はEOFまでをコードとして扱う)
            blocks.append({"type": "codeblock", "text": "\n".join(code_lines)})
            continue

        # 表: 現在行が | を含み、次行が区切り行(---|---)
        if "|" in stripped and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1].strip()):
            header = [parse_inline(c) for c in split_table_row(stripped)]
            i += 2
            rows = []
            while i < n and "|" in lines[i].strip() and lines[i].strip() != "":
                rows.append([parse_inline(c) for c in split_table_row(lines[i].strip())])
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "quote", "runs": parse_inline(" ".join(quote_lines))})
            continue

        m_ul = UL_RE.match(stripped)
        m_ol = OL_RE.match(stripped)
        if m_ul or m_ol:
            ordered = bool(m_ol)
            items = []
            while i < n:
                s = lines[i].strip()
                mm = OL_RE.match(s) if ordered else UL_RE.match(s)
                if not mm:
                    break
                items.append(parse_inline(mm.group(1).strip()))
                i += 1
            blocks.append({"type": "list", "ordered": ordered, "items": items})
            continue

        # 段落: 空行/他ブロック開始まで連結
        para_lines = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if nxt == "" or HEADING_RE.match(nxt) or HR_RE.match(nxt) or nxt.startswith(">") \
                    or nxt.startswith("```") \
                    or UL_RE.match(nxt) or OL_RE.match(nxt) or ("|" in nxt and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1].strip())):
                break
            para_lines.append(nxt)
            i += 1
        blocks.append({"type": "para", "runs": parse_inline(" ".join(para_lines))})

    return blocks


# ---------------------------------------------------------------------------
# 2. タイプ別の後処理(見出し番号・配布先欄強調)
# ---------------------------------------------------------------------------

def apply_yoken_numbering(blocks):
    """yokenタイプ: h2〜h4に章番号を振る(h1=タイトルは対象外)。"""
    counters = [0, 0, 0]  # h2, h3, h4
    for b in blocks:
        if b["type"] != "heading":
            continue
        level = b["level"]
        if level == 1:
            continue
        idx = level - 2  # 0,1,2
        counters[idx] += 1
        for j in range(idx + 1, 3):
            counters[j] = 0
        number = ".".join(str(c) for c in counters[: idx + 1] if c) + "."
        if b["runs"]:
            b["runs"][0] = dict(b["runs"][0])
            b["runs"][0]["text"] = f"{number} {b['runs'][0]['text']}"
        else:
            b["runs"] = [{"text": number, "bold": False, "code": False}]
    return blocks


def is_dist_row(row_runs) -> bool:
    if not row_runs:
        return False
    first_text = "".join(r["text"] for r in row_runs[0]).strip()
    return first_text in ("配布先", "宛先")


def is_meta_table(header) -> bool:
    """議事録冒頭の項目/内容キーバリュー表(配布先欄などを含む)かどうか。
    決定事項表・宿題表など見出し付きの通常表と区別し、見出し非表示化を誤適用しない。"""
    header_text = ["".join(r["text"] for r in cell).strip() for cell in header]
    return header_text[:2] == ["項目", "内容"]


# ---------------------------------------------------------------------------
# 3. HTML レンダラ(標準ライブラリのみ)
# ---------------------------------------------------------------------------

def esc(s: str) -> str:
    return html_lib.escape(s, quote=False)


def runs_to_html(runs) -> str:
    out = []
    for r in runs:
        t = esc(r["text"])
        if r["code"]:
            t = f"<code>{t}</code>"
        if r["bold"]:
            t = f"<strong>{t}</strong>"
        out.append(t)
    return "".join(out)


TYPE_LABELS = {"gijiroku": "議事録", "yoken": "要件定義", "kaito": "回答"}


def blocks_to_html_body(blocks, doc_type: str) -> str:
    parts = []
    for b in blocks:
        if b["type"] == "heading":
            level = b["level"]
            parts.append(f"<h{level}>{runs_to_html(b['runs'])}</h{level}>")
        elif b["type"] == "para":
            parts.append(f"<p>{runs_to_html(b['runs'])}</p>")
        elif b["type"] == "hr":
            parts.append("<hr>")
        elif b["type"] == "codeblock":
            parts.append(f"<pre><code>{esc(b['text'])}</code></pre>")
        elif b["type"] == "quote":
            parts.append(f"<blockquote>{runs_to_html(b['runs'])}</blockquote>")
        elif b["type"] == "list":
            tag = "ol" if b["ordered"] else "ul"
            items = "".join(f"<li>{runs_to_html(it)}</li>" for it in b["items"])
            parts.append(f"<{tag}>{items}</{tag}>")
        elif b["type"] == "table":
            table_class = "meta-table" if (doc_type == "gijiroku" and is_meta_table(b["header"])) else ""
            rows_html = []
            rows_html.append(
                "<tr>" + "".join(f"<th>{runs_to_html(c)}</th>" for c in b["header"]) + "</tr>"
            )
            for row in b["rows"]:
                tr_class = ' class="dist-row"' if doc_type == "gijiroku" and is_dist_row(row) else ""
                rows_html.append(
                    f"<tr{tr_class}>" + "".join(f"<td>{runs_to_html(c)}</td>" for c in row) + "</tr>"
                )
            parts.append(f'<table class="{table_class}">' + "".join(rows_html) + "</table>")
    return "\n".join(parts)


def build_html(blocks, doc_type: str, title: str) -> str:
    page_size = "A4 landscape" if doc_type == "kaito" else "A4"
    body = blocks_to_html_body(blocks, doc_type)

    css = f"""
@page {{ size: {page_size}; margin: 18mm 16mm; }}
* {{ box-sizing: border-box; }}
body {{
  font-family: "Yu Gothic", "游ゴシック", "Hiragino Kaku Gothic ProN", "Meiryo", sans-serif;
  color: #1a1a1a;
  line-height: 1.7;
  max-width: 900px;
  margin: 0 auto;
  padding: 24px;
}}
h1 {{ font-size: 20px; border-bottom: 3px solid #333; padding-bottom: 6px; }}
h2 {{ font-size: 17px; border-left: 6px solid #555; padding-left: 8px; margin-top: 28px; }}
h3 {{ font-size: 15px; margin-top: 20px; }}
h4 {{ font-size: 13.5px; margin-top: 16px; }}
table {{ border-collapse: collapse; width: 100%; margin: 12px 0 20px; font-size: 13px; }}
th, td {{ border: 1px solid #444; padding: 5px 9px; text-align: left; vertical-align: top; }}
th {{ background: #eaeaea; }}
blockquote {{ border-left: 4px solid #999; margin: 12px 0; padding: 6px 12px; color: #444; background: #f7f7f7; }}
hr {{ border: none; border-top: 1px solid #999; margin: 20px 0; }}
code {{ background: #f0f0f0; padding: 1px 5px; font-family: "Courier New", monospace; font-size: 0.95em; }}
pre {{ background: #f5f5f5; border: 1px solid #ddd; padding: 10px 12px; overflow-x: auto; font-family: "MS Gothic", "ＭＳ ゴシック", "Courier New", monospace; font-size: 0.85em; line-height: 1.4; }}
pre code {{ background: none; padding: 0; }}
ul, ol {{ padding-left: 24px; margin: 8px 0; }}
li {{ margin: 3px 0; }}
.meta-table th {{ display: none; }}
.meta-table td:first-child {{ background: #eee; font-weight: bold; width: 18%; }}
tr.dist-row td {{ background: #fff2cc; font-weight: bold; }}
.footer-note {{ margin-top: 32px; padding: 10px 12px; background: #fff8e1; border: 1px solid #e0c060; font-size: 13px; }}
@media print {{
  body {{ padding: 0; }}
  .footer-note {{ break-inside: avoid; }}
}}
"""
    return f"""<!doctype html>
<html lang="ja">
<head>
<meta charset="utf-8">
<title>{esc(title)}</title>
<style>{css}</style>
</head>
<body>
{body}
<div class="footer-note">⚠️ 正本昇格前に目視確認(docs/02 M-1)。docx手動修正後はdocxが正本、mdには⚠️注記を追記すること(M-2)。</div>
</body>
</html>
"""


# ---------------------------------------------------------------------------
# 4. docx レンダラ(python-docxがある場合のみ)
# ---------------------------------------------------------------------------

def build_docx(blocks, doc_type: str, title: str, out_path: Path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    FONT_NAME = "游ゴシック"

    def set_east_asian_font(run):
        run.font.name = FONT_NAME
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), FONT_NAME)

    def add_runs(paragraph, runs):
        if not runs:
            paragraph.add_run("")
            return
        for r in runs:
            run = paragraph.add_run(r["text"])
            run.bold = r["bold"]
            if r["code"]:
                run.font.name = "Courier New"
            else:
                set_east_asian_font(run)

    def shade_cell(cell, color_hex: str):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color_hex)
        tcPr.append(shd)

    def add_hr(document):
        p = document.add_paragraph()
        p_fmt = p.paragraph_format
        p_fmt.space_before = Pt(2)
        p_fmt.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "999999")
        pBdr.append(bottom)
        pPr.append(pBdr)

    document = Document()

    # 既定フォント
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10.5)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # A4・kaitoは横向き可
    section = document.sections[0]
    if doc_type == "kaito":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(1.8)

    for b in blocks:
        if b["type"] == "heading":
            level = min(b["level"], 4)
            p = document.add_heading("", level=level)
            add_runs(p, b["runs"])
            for run in p.runs:
                if not run.font.name or run.font.name == FONT_NAME:
                    set_east_asian_font(run)
        elif b["type"] == "para":
            p = document.add_paragraph()
            add_runs(p, b["runs"])
        elif b["type"] == "hr":
            add_hr(document)
        elif b["type"] == "codeblock":
            # 等幅フォントの段落として1行ずつ追加(整形を保つ。日本語はMSゴシックをeastAsia指定)
            for code_line in b["text"].split("\n"):
                p = document.add_paragraph()
                run = p.add_run(code_line if code_line else " ")
                run.font.name = "MS Gothic"
                run.font.size = Pt(9)
                rPr = run._element.get_or_add_rPr()
                rFonts = rPr.find(qn("w:rFonts"))
                if rFonts is None:
                    rFonts = OxmlElement("w:rFonts")
                    rPr.append(rFonts)
                rFonts.set(qn("w:eastAsia"), "ＭＳ ゴシック")
        elif b["type"] == "quote":
            p = document.add_paragraph(style="Intense Quote") if "Intense Quote" in [
                s.name for s in document.styles
            ] else document.add_paragraph()
            add_runs(p, b["runs"])
            p.paragraph_format.left_indent = Cm(0.8)
        elif b["type"] == "list":
            style = "List Number" if b["ordered"] else "List Bullet"
            for item_runs in b["items"]:
                p = document.add_paragraph(style=style)
                add_runs(p, item_runs)
        elif b["type"] == "table":
            ncols = len(b["header"])
            table = document.add_table(rows=1, cols=max(ncols, 1))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = table.rows[0].cells
            for j, cell_runs in enumerate(b["header"]):
                p = hdr_cells[j].paragraphs[0]
                add_runs(p, cell_runs)
                for run in p.runs:
                    run.bold = True
                if doc_type == "gijiroku":
                    shade_cell(hdr_cells[j], "D9D9D9")
            for row in b["rows"]:
                cells = table.add_row().cells
                dist = doc_type == "gijiroku" and is_dist_row(row)
                for j, cell_runs in enumerate(row):
                    if j >= len(cells):
                        break
                    p = cells[j].paragraphs[0]
                    add_runs(p, cell_runs)
                    if dist:
                        shade_cell(cells[j], "FFF2CC")
                        for run in p.runs:
                            run.bold = True

    footer_p = document.add_paragraph()
    footer_run = footer_p.add_run(
        "⚠️ 正本昇格前に目視確認(docs/02 M-1)。docx手動修正後はdocxが正本、mdには⚠️注記を追記すること(M-2)。"
    )
    footer_run.italic = True
    set_east_asian_font(footer_run)

    document.save(str(out_path))


# ---------------------------------------------------------------------------
# 5. 出力パスの決定・既存正本の退避
# ---------------------------------------------------------------------------

def reserve_output_path(path: Path) -> Path:
    """既存出力があれば上書きせず _2(以降空き番号)を付けて退避し、元の名前を空ける。"""
    if not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    n = 2
    while True:
        candidate = path.with_name(f"{stem}_{n}{suffix}")
        if not candidate.exists():
            break
        n += 1
    path.rename(candidate)
    print(f"  ↳ 既存出力を退避: {path.name} → {candidate.name}(正本保護のため上書きしません)")
    return path


# ---------------------------------------------------------------------------
# 6. CLI
# ---------------------------------------------------------------------------

def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="md2office.py",
        description="md→docx/html 統一変換スクリプト(業務OS 候補13)。読み取り専用・正本上書きなし。",
    )
    p.add_argument("input", help="入力.md のパス")
    p.add_argument(
        "--type",
        dest="doc_type",
        default=None,
        help=f"文書タイプ: {'|'.join(VALID_TYPES)}",
    )
    p.add_argument(
        "--out",
        default="docx,html",
        help="出力形式(カンマ区切り)。既定: docx,html",
    )
    p.add_argument(
        "--outdir",
        default=None,
        help="出力先ディレクトリ。省略時は入力ファイルと同じディレクトリ",
    )
    return p


def run(argv) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    # --type 検証(argparse choicesを使わず、常に終了コード0でメッセージを出す)
    if args.doc_type is None:
        print("❌ エラー: --type が指定されていません(gijiroku|yoken|kaito のいずれかを指定)", file=sys.stderr)
        return 0
    if args.doc_type not in VALID_TYPES:
        print(
            f"❌ エラー: --type '{args.doc_type}' は不正です。使えるのは次のいずれか: {', '.join(VALID_TYPES)}",
            file=sys.stderr,
        )
        return 0

    in_path = Path(args.input)
    if not in_path.is_file():
        print(f"❌ エラー: 入力ファイルが見つかりません: {in_path}", file=sys.stderr)
        return 0

    requested_outs = [o.strip().lower() for o in args.out.split(",") if o.strip()]
    outs = [o for o in requested_outs if o in VALID_OUTS]
    bad_outs = [o for o in requested_outs if o not in VALID_OUTS]
    if bad_outs:
        print(f"⚠️ 警告: 未対応の --out 指定を無視します: {', '.join(bad_outs)}(対応: {', '.join(VALID_OUTS)})", file=sys.stderr)
    if not outs:
        print("❌ エラー: 有効な --out 指定がありません(docx,html のいずれかを含めてください)", file=sys.stderr)
        return 0

    outdir = Path(args.outdir) if args.outdir else in_path.parent
    try:
        outdir.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        print(f"❌ エラー: 出力先ディレクトリを作成できません: {outdir}({e})", file=sys.stderr)
        return 0

    text = in_path.read_text(encoding="utf-8")
    blocks = parse_markdown(text)
    if args.doc_type == "yoken":
        blocks = apply_yoken_numbering(blocks)

    title = in_path.stem
    stem = in_path.stem

    print(f"入力: {in_path}(type={args.doc_type})")

    if "html" in outs:
        html_path = outdir / f"{stem}.html"
        html_path = reserve_output_path(html_path)
        html_text = build_html(blocks, args.doc_type, title)
        html_path.write_text(html_text, encoding="utf-8")
        print(f"  ✅ html出力: {html_path}")

    if "docx" in outs:
        docx_path = outdir / f"{stem}.docx"
        try:
            import docx  # noqa: F401  (存在確認のみ)
        except ImportError:
            print(
                "⚠️ python-docx が見つかりません。docx出力をスキップしhtmlのみで継続します。\n"
                "   導入するには: pip install python-docx",
                file=sys.stderr,
            )
        else:
            docx_path = reserve_output_path(docx_path)
            try:
                build_docx(blocks, args.doc_type, title, docx_path)
                print(f"  ✅ docx出力: {docx_path}")
            except Exception as e:  # noqa: BLE001 生成失敗でも落とさない
                print(f"❌ docx生成に失敗しました: {e}", file=sys.stderr)

    print("⚠️ 正本昇格前に目視確認(docs/02 M-1)")
    return 0


def main():
    try:
        code = run(sys.argv[1:])
    except SystemExit:
        # argparseの使用法エラー等も含め、レポートツールとして終了コードは常に0に統一
        code = 0
    except Exception as e:  # noqa: BLE001
        print(f"❌ 予期しないエラー: {e}", file=sys.stderr)
        code = 0
    sys.exit(code if code is not None else 0)


if __name__ == "__main__":
    main()
