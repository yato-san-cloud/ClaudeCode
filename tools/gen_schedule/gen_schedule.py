#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""gen_schedule.py — スケジュール3型(WBS/支援者/移管)をデータ1式から自動生成する。

使い方:
    python gen_schedule.py tasks.csv [--out DIR] [--today YYYY-MM-DD]   # WBS型のみ
    python gen_schedule.py プロジェクトフォルダ [--out DIR] [--today ...]  # 3型連動

プロジェクトフォルダの構成(ファイル名で自動発見。tasks.csv 以外は任意):
    tasks.csv        … WBS型タスク表(必須)
    assignments.csv  … 支援者割当表(リソース型)
    lots.csv         … 移管ロット定義 + quota.csv … 日別計画/実績(数量型)

列仕様は README.md を参照。3表は「関連タスクID」で tasks.csv に紐づき、
期間のはみ出し・数量の不整合は実行時に整合チェックとして警告する。

出力(--out 配下):
    <名前>_ガント.xlsx / .html   … WBS週次ガント(□■☆◆ / 担当フィルタ・今日線)
    支援者_マトリクス.xlsx / .html … 人×日付の割当マトリクス
    移管_マトリクス.xlsx          … ロット×日別の計画/実績
    移管_バーンダウン.html        … 累積計画vs実績の消化曲線
    スケジュール_最新.md          … 3型統合スナップショット(エージェント用ナレッジ)
    変更履歴.md                  … 3表まとめて前回実行とのdiffを自動追記
"""
import argparse, csv, datetime as dt, html, json, pathlib, sys

STATES = ("未着手", "着手", "完了")


def parse_date(s):
    s = str(s).strip().split(" ")[0]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"日付を解釈できない: {s!r}")


def read_table(path):
    path = pathlib.Path(path)
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        header = [str(v).strip() if v else "" for v in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def g(r, k):
    return str(r.get(k) or "").strip()


# ---------- 読み込み ----------
def load_tasks(path):
    tasks, errors = [], []
    for i, r in enumerate(read_table(path), start=2):
        if not g(r, "ID"):
            errors.append(f"{i}行目: ID が空"); continue
        try:
            t = {"id": g(r, "ID"), "l1": g(r, "大項目"), "l2": g(r, "中項目"), "l3": g(r, "小項目"),
                 "own": g(r, "担当正"), "sub": g(r, "担当副"),
                 "start": parse_date(g(r, "開始日")), "end": parse_date(g(r, "終了日")),
                 "state": g(r, "状態") or "未着手", "ms": bool(g(r, "マイルストーン")),
                 "dep": g(r, "先行ID"), "note": g(r, "進捗メモ"), "memo": g(r, "備考")}
        except ValueError as e:
            errors.append(f"{i}行目({g(r,'ID')}): {e}"); continue
        if t["state"] not in STATES:
            errors.append(f"{i}行目({t['id']}): 状態は {'/'.join(STATES)} のみ"); continue
        if t["end"] < t["start"]:
            errors.append(f"{i}行目({t['id']}): 終了日が開始日より前"); continue
        tasks.append(t)
    if errors:
        sys.exit("tasks 入力エラー:\n  " + "\n  ".join(errors))
    return tasks


def load_assignments(path):
    rows, errors = [], []
    for i, r in enumerate(read_table(path), start=2):
        try:
            rows.append({"date": parse_date(g(r, "日付")), "who": g(r, "氏名"), "org": g(r, "所属"),
                         "place": g(r, "場所"), "time": g(r, "時間帯"), "work": g(r, "作業内容"),
                         "task": g(r, "関連タスクID"), "memo": g(r, "備考")})
        except ValueError as e:
            errors.append(f"{i}行目: {e}")
    if errors:
        sys.exit("assignments 入力エラー:\n  " + "\n  ".join(errors))
    return [r for r in rows if r["who"]]


def load_quota(lots_path, quota_path):
    lots, errors = {}, []
    for i, r in enumerate(read_table(lots_path), start=2):
        key = (g(r, "移管元"), g(r, "荷姿"), g(r, "頻度"))
        try:
            lots[key] = {"total": int(float(g(r, "総量") or 0)), "task": g(r, "関連タスクID"),
                         "memo": g(r, "備考")}
        except ValueError:
            errors.append(f"lots {i}行目: 総量が数値でない")
    quota = []
    for i, r in enumerate(read_table(quota_path), start=2):
        key = (g(r, "移管元"), g(r, "荷姿"), g(r, "頻度"))
        try:
            quota.append({"key": key, "date": parse_date(g(r, "日付")),
                          "plan": int(float(g(r, "計画") or 0)),
                          "act": None if g(r, "実績") == "" else int(float(g(r, "実績")))})
        except ValueError as e:
            errors.append(f"quota {i}行目: {e}")
        if key not in lots:
            errors.append(f"quota {i}行目: lots.csv に無いロット {key}")
    if errors:
        sys.exit("移管 入力エラー:\n  " + "\n  ".join(errors))
    return lots, quota


# ---------- 整合チェック(連動の中核) ----------
def cross_check(tasks, assigns, lots, quota):
    warn = []
    tmap = {t["id"]: t for t in tasks}
    for a in assigns:
        if a["task"]:
            t = tmap.get(a["task"])
            if not t:
                warn.append(f"支援者: {a['date']} {a['who']} の関連タスク {a['task']} が tasks に無い")
            elif not (t["start"] <= a["date"] <= t["end"]):
                warn.append(f"支援者: {a['date']} {a['who']} の割当が {a['task']}"
                            f"({t['start'].strftime('%m/%d')}〜{t['end'].strftime('%m/%d')})の期間外")
    for key, lot in lots.items():
        if lot["task"] and lot["task"] not in tmap:
            warn.append(f"移管: ロット{'/'.join(key)} の関連タスク {lot['task']} が tasks に無い")
        plan_sum = sum(q["plan"] for q in quota if q["key"] == key)
        if plan_sum != lot["total"]:
            warn.append(f"移管: ロット{'/'.join(key)} 計画合計{plan_sum} ≠ 総量{lot['total']}(差{plan_sum - lot['total']:+d})")
        t = tmap.get(lot["task"]) if lot["task"] else None
        if t:
            outside = [q for q in quota if q["key"] == key and not (t["start"] <= q["date"] <= t["end"])]
            if outside:
                warn.append(f"移管: ロット{'/'.join(key)} の計画{len(outside)}日分が {lot['task']} の期間外")
    return warn


def monday(d):
    return d - dt.timedelta(days=d.weekday())


def week_axis(tasks):
    lo = monday(min(t["start"] for t in tasks))
    hi = monday(max(t["end"] for t in tasks))
    weeks = []
    w = lo
    while w <= hi:
        weeks.append(w); w += dt.timedelta(days=7)
    return weeks


THIN = None
def _styles():
    from openpyxl.styles import Border, Side
    global THIN
    THIN = Border(*[Side(style="thin", color="BBBBBB")] * 4)


# ---------- WBS型: Excel 週次ガント ----------
def write_xlsx(tasks, weeks, today, out, name):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    _styles()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "ガント"
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    month_fill = PatternFill("solid", fgColor="EEEEEE")
    today_fill = PatternFill("solid", fgColor="FFF2CC")
    cols = ["ID", "大項目", "中項目", "小項目", "担当正", "担当副", "開始", "終了", "状態", "進捗メモ"]
    ws.cell(1, 1, f"{name}  (生成: {today}  □=予定 ■=実施 ☆=完了 ◆=ﾏｲﾙｽﾄｰﾝ)").font = Font(bold=True)
    for j, h in enumerate(cols, 1):
        c = ws.cell(3, j, h); c.fill = head_fill; c.border = THIN; c.font = Font(bold=True, size=9)
    base = len(cols)
    for k, w in enumerate(weeks):
        c1 = ws.cell(2, base + 1 + k)
        if k == 0 or w.month != weeks[k - 1].month:
            c1.value = f"{w.year}/{w.month}月"; c1.font = Font(bold=True, size=9); c1.fill = month_fill
        c2 = ws.cell(3, base + 1 + k, f"{w.month}/{w.day}")
        c2.font = Font(size=8); c2.fill = head_fill; c2.border = THIN
        c2.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c2.column_letter].width = 4.5
    widths = {"A": 5, "B": 11, "C": 11, "D": 30, "E": 8, "F": 8, "G": 9, "H": 9, "I": 6, "J": 24}
    for col, wd in widths.items():
        ws.column_dimensions[col].width = wd
    prev_l1 = prev_l2 = None
    for i, t in enumerate(tasks, start=4):
        vals = [t["id"], t["l1"] if t["l1"] != prev_l1 else "", t["l2"] if t["l2"] != prev_l2 else "",
                t["l3"], t["own"], t["sub"], t["start"].strftime("%m/%d"), t["end"].strftime("%m/%d"),
                t["state"], t["note"]]
        prev_l1, prev_l2 = t["l1"], t["l2"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.border = THIN; c.font = Font(size=9)
            c.alignment = Alignment(vertical="center", wrap_text=(j in (4, 10)))
        for k, w in enumerate(weeks):
            c = ws.cell(i, base + 1 + k); c.border = THIN
            c.alignment = Alignment(horizontal="center"); c.font = Font(size=9)
            we = w + dt.timedelta(days=6)
            in_span = not (we < t["start"] or w > t["end"])
            mark = ""
            if in_span:
                if t["state"] == "完了":
                    mark = "■"
                elif t["state"] == "着手":
                    mark = "■" if we <= today else "□"
                else:
                    mark = "□"
            if t["ms"] and w <= t["end"] <= we:
                mark = "◆"
            elif t["state"] == "完了" and w <= t["end"] <= we:
                mark = "☆"
            c.value = mark
            if w <= today <= we:
                c.fill = today_fill
    ws.freeze_panes = ws.cell(4, base + 1)
    wb.save(out / f"{name}_ガント.xlsx")


# ---------- WBS型: HTML ガント ----------
def write_html(tasks, weeks, today, out, name):
    day0 = weeks[0]
    days = (weeks[-1] + dt.timedelta(days=7) - day0).days
    px = 4.2
    W = int(days * px) + 340
    owners = sorted({t["own"] for t in tasks if t["own"]})
    rows_html, y = [], 0
    prev_l1 = None
    for t in tasks:
        if t["l1"] != prev_l1:
            rows_html.append(f'<div class="grp" style="top:{y}px">{html.escape(t["l1"])}</div>')
            y += 26; prev_l1 = t["l1"]
        x = (t["start"] - day0).days * px + 320
        wpx = max(((t["end"] - t["start"]).days + 1) * px, 6)
        cls = {"未着手": "todo", "着手": "doing", "完了": "done"}[t["state"]]
        late = t["state"] != "完了" and t["end"] < today
        tip = html.escape(f'{t["id"]} {t["l3"]} / {t["own"]}(副:{t["sub"] or "-"}) '
                          f'{t["start"]}〜{t["end"]} [{t["state"]}]'
                          + (f' 先行:{t["dep"]}' if t["dep"] else "") + (f' | {t["note"]}' if t["note"] else ""))
        bar = (f'<div class="ms" style="left:{x + wpx - 7:.0f}px;top:{y + 4}px" title="{tip}">◆</div>'
               if t["ms"] else
               f'<div class="bar {cls}{" late" if late else ""}" style="left:{x:.0f}px;top:{y + 5}px;width:{wpx:.0f}px" title="{tip}"></div>')
        rows_html.append(
            f'<div class="row" data-own="{html.escape(t["own"])}" style="top:{y}px">'
            f'<div class="nm" title="{tip}">{html.escape(t["id"])} {html.escape(t["l3"] or t["l2"])}'
            f'<span class="ow">{html.escape(t["own"])}</span></div>{bar}</div>')
        y += 24
    H = y + 60
    grid = []
    for w in weeks:
        gx = (w - day0).days * px + 320
        grid.append(f'<div class="gl" style="left:{gx:.0f}px;height:{H}px"></div>')
        if w.day <= 7 or w == weeks[0]:
            grid.append(f'<div class="glab" style="left:{gx:.0f}px">{w.year}/{w.month}月</div>')
    tx = (today - day0).days * px + 320
    opts = "".join(f'<option>{html.escape(o)}</option>' for o in owners)
    doc = f"""<!doctype html><html lang="ja"><head><meta charset="utf-8">
<title>{html.escape(name)} ガント</title><style>
body{{font-family:"Segoe UI","Yu Gothic UI",sans-serif;margin:12px;color:#1a1a1a}}
h1{{font-size:16px}} .legend span{{margin-right:14px;font-size:12px}}
#wrap{{position:relative;width:{W}px;height:{H}px;border:1px solid #ccc;overflow:hidden;background:#fff}}
.row{{position:absolute;left:0;width:100%;height:24px}}
.nm{{position:absolute;left:0;width:312px;font-size:11px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;padding:4px 4px 0 12px;background:#fff;z-index:2}}
.ow{{color:#888;margin-left:6px;font-size:10px}}
.grp{{position:absolute;left:0;width:100%;height:26px;background:#eef4fb;font-weight:bold;font-size:12px;padding:4px 0 0 4px;z-index:3}}
.bar{{position:absolute;height:12px;border-radius:3px;z-index:1}}
.bar.todo{{background:#c7d8ee}} .bar.doing{{background:#2f7bd0}} .bar.done{{background:#9fc99f}}
.bar.late{{outline:2px solid #d03020}}
.ms{{position:absolute;color:#d08000;font-size:13px;z-index:2}}
.gl{{position:absolute;top:0;width:1px;background:#eee}}
.glab{{position:absolute;top:2px;font-size:10px;color:#666}}
#today{{position:absolute;top:0;width:2px;background:#d03020;height:{H}px;z-index:4}}
select{{font-size:12px}}
</style></head><body>
<h1>{html.escape(name)} <small>(生成: {today})</small></h1>
<p class="legend"><span style="color:#c7d8ee">■</span>未着手 <span style="color:#2f7bd0">■</span>着手
<span style="color:#9fc99f">■</span>完了 <span style="color:#d03020">□</span>遅延(赤枠) <span style="color:#d08000">◆</span>ﾏｲﾙｽﾄｰﾝ
　担当フィルタ: <select id="f" onchange="flt()"><option>全員</option>{opts}</select></p>
<div id="wrap">{"".join(grid)}<div id="today" style="left:{tx:.0f}px"></div>{"".join(rows_html)}</div>
<script>
function flt(){{const v=document.getElementById('f').value;
document.querySelectorAll('.row').forEach(r=>{{r.style.opacity=(v==='全員'||r.dataset.own===v)?1:0.12;}});}}
</script></body></html>"""
    (out / f"{name}_ガント.html").write_text(doc, encoding="utf-8")


# ---------- リソース型: 支援者マトリクス ----------
PLACE_COLORS = ["E2EFDA", "DDEBF7", "FCE4D6", "E4DFEC", "FFF2CC", "D9E1F2"]


def write_assign_outputs(assigns, today, out):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    _styles()
    dates = sorted({a["date"] for a in assigns})
    people = sorted({(a["org"], a["who"]) for a in assigns})
    places = sorted({a["place"] for a in assigns if a["place"]})
    color_of = {p: PLACE_COLORS[i % len(PLACE_COLORS)] for i, p in enumerate(places)}
    amap = {}
    for a in assigns:
        amap.setdefault((a["who"], a["date"]), []).append(a)
    # xlsx
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "支援者マトリクス"
    ws.cell(1, 1, f"支援者スケジュール (生成: {today})").font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for j, h in enumerate(["所属", "氏名"], 1):
        c = ws.cell(2, j, h); c.fill = head_fill; c.border = THIN; c.font = Font(bold=True, size=9)
        ws.column_dimensions[c.column_letter].width = 10
    for k, d in enumerate(dates):
        c = ws.cell(2, 3 + k, f"{d.month}/{d.day}({'月火水木金土日'[d.weekday()]})")
        c.fill = head_fill; c.border = THIN; c.font = Font(bold=True, size=8)
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = 14
    for i, (org, who) in enumerate(people, start=3):
        ws.cell(i, 1, org).font = Font(size=9); ws.cell(i, 1).border = THIN
        ws.cell(i, 2, who).font = Font(size=9); ws.cell(i, 2).border = THIN
        for k, d in enumerate(dates):
            c = ws.cell(i, 3 + k); c.border = THIN; c.font = Font(size=8)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            items = amap.get((who, d), [])
            if items:
                c.value = "\n".join(f'{x["place"]}{"(" + x["time"] + ")" if x["time"] else ""} {x["work"]}'.strip()
                                    for x in items)
                if items[0]["place"] in color_of:
                    c.fill = PatternFill("solid", fgColor=color_of[items[0]["place"]])
    ws.cell(len(people) + 4, 1, "凡例: セル色=場所 / 割当日数計").font = Font(size=9, bold=True)
    for i, (org, who) in enumerate(people):
        n = sum(1 for d in dates if (who, d) in amap)
        ws.cell(len(people) + 5 + i, 1, f"{who}: {n}日").font = Font(size=9)
    ws.freeze_panes = "C3"
    wb.save(out / "支援者_マトリクス.xlsx")
    # html
    th = "".join(f'<th>{d.month}/{d.day}<br>({"月火水木金土日"[d.weekday()]})</th>' for d in dates)
    trs = []
    for org, who in people:
        tds = []
        for d in dates:
            items = amap.get((who, d), [])
            if items:
                bg = f'#{color_of.get(items[0]["place"], "EEEEEE")}'
                cell = "<br>".join(html.escape(f'{x["place"]} {x["work"]}'.strip()) for x in items)
                tip = html.escape(" / ".join(f'{x["place"]}{x["time"] and "(" + x["time"] + ")"} {x["work"]}'
                                             + (f' [{x["task"]}]' if x["task"] else "") for x in items))
                tds.append(f'<td style="background:{bg}" title="{tip}">{cell}</td>')
            else:
                tds.append("<td></td>")
        n = sum(1 for d in dates if (who, d) in amap)
        trs.append(f'<tr><td class="p">{html.escape(org)}</td><td class="p"><b>{html.escape(who)}</b> '
                   f'<span class="n">{n}日</span></td>{"".join(tds)}</tr>')
    doc = f"""<!doctype html><html lang="ja"><head><meta charset="utf-8"><title>支援者スケジュール</title><style>
body{{font-family:"Segoe UI","Yu Gothic UI",sans-serif;margin:12px;color:#1a1a1a;font-size:12px}}
table{{border-collapse:collapse}} th,td{{border:1px solid #bbb;padding:3px 6px;font-size:11px;min-width:90px;vertical-align:top}}
th{{background:#ddebf7}} td.p{{background:#fff;white-space:nowrap;min-width:40px}} .n{{color:#888;font-size:10px}}
</style></head><body><h1 style="font-size:16px">支援者スケジュール <small>(生成: {today})</small></h1>
<table><tr><th>所属</th><th>氏名</th>{th}</tr>{"".join(trs)}</table></body></html>"""
    (out / "支援者_マトリクス.html").write_text(doc, encoding="utf-8")


# ---------- 数量型: 移管マトリクス + バーンダウン ----------
def write_quota_outputs(lots, quota, today, out):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    _styles()
    dates = sorted({q["date"] for q in quota})
    qmap = {(q["key"], q["date"]): q for q in quota}
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "移管マトリクス"
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    done_fill = PatternFill("solid", fgColor="E2EFDA")
    short_fill = PatternFill("solid", fgColor="FCE4D6")
    ws.cell(1, 1, f"移管スケジュール (生成: {today}  セル=計画PL、実績記入済は 実績/計画 表示。"
                  "緑=計画達成 橙=未達)").font = Font(bold=True)
    heads = ["移管元", "荷姿", "頻度", "総量", "計画計", "実績計", "消化率"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(2, j, h); c.fill = head_fill; c.border = THIN; c.font = Font(bold=True, size=9)
        ws.column_dimensions[c.column_letter].width = 8
    for k, d in enumerate(dates):
        c = ws.cell(2, len(heads) + 1 + k, f"{d.month}/{d.day}")
        c.fill = head_fill; c.border = THIN; c.font = Font(bold=True, size=8)
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = 7
    for i, (key, lot) in enumerate(sorted(lots.items()), start=3):
        plan_sum = sum(q["plan"] for q in quota if q["key"] == key)
        act_sum = sum(q["act"] or 0 for q in quota if q["key"] == key)
        rate = f"{act_sum / lot['total'] * 100:.0f}%" if lot["total"] else "-"
        for j, v in enumerate([*key, lot["total"], plan_sum, act_sum, rate], 1):
            c = ws.cell(i, j, v); c.border = THIN; c.font = Font(size=9)
        for k, d in enumerate(dates):
            c = ws.cell(i, len(heads) + 1 + k); c.border = THIN; c.font = Font(size=8)
            c.alignment = Alignment(horizontal="center")
            q = qmap.get((key, d))
            if q and (q["plan"] or q["act"] is not None):
                if q["act"] is None:
                    c.value = q["plan"] or ""
                else:
                    c.value = f'{q["act"]}/{q["plan"]}'
                    c.fill = done_fill if q["act"] >= q["plan"] else short_fill
    ws.freeze_panes = ws.cell(3, len(heads) + 1)
    wb.save(out / "移管_マトリクス.xlsx")
    # バーンダウン(累積 計画 vs 実績)
    total = sum(l["total"] for l in lots.values())
    cum_p, cum_a, cp, ca = [], [], 0, 0
    has_act_dates = [q["date"] for q in quota if q["act"] is not None]
    last_act = max(has_act_dates) if has_act_dates else None
    for d in dates:
        cp += sum(q["plan"] for q in quota if q["date"] == d)
        cum_p.append(cp)
        ca += sum(q["act"] or 0 for q in quota if q["date"] == d)
        cum_a.append(ca if last_act and d <= last_act else None)
    Wp, Hp, pad = 720, 300, 46
    n = max(len(dates) - 1, 1)
    mx = max(total, cp) or 1
    def pt(i, v):
        return f"{pad + i * (Wp - 2 * pad) / n:.1f},{Hp - pad - v * (Hp - 2 * pad) / mx:.1f}"
    pline = " ".join(pt(i, v) for i, v in enumerate(cum_p))
    aline = " ".join(pt(i, v) for i, v in enumerate(cum_a) if v is not None)
    xlabels = "".join(f'<text x="{pad + i * (Wp - 2 * pad) / n:.0f}" y="{Hp - pad + 14}" font-size="9" '
                      f'text-anchor="middle">{d.month}/{d.day}</text>'
                      for i, d in enumerate(dates) if i % max(1, len(dates) // 12) == 0)
    prog = f"{ca}/{total} PL ({ca / total * 100:.0f}%)" if total else "-"
    doc = f"""<!doctype html><html lang="ja"><head><meta charset="utf-8"><title>移管バーンダウン</title><style>
body{{font-family:"Segoe UI","Yu Gothic UI",sans-serif;margin:12px;color:#1a1a1a}}</style></head><body>
<h1 style="font-size:16px">移管消化曲線 <small>(生成: {today} / 実績 {prog})</small></h1>
<svg width="{Wp}" height="{Hp}" style="border:1px solid #ccc;background:#fff">
<line x1="{pad}" y1="{Hp - pad}" x2="{Wp - pad}" y2="{Hp - pad}" stroke="#999"/>
<line x1="{pad}" y1="{pad}" x2="{pad}" y2="{Hp - pad}" stroke="#999"/>
<line x1="{pad}" y1="{Hp - pad - total * (Hp - 2 * pad) / mx:.1f}" x2="{Wp - pad}" y2="{Hp - pad - total * (Hp - 2 * pad) / mx:.1f}" stroke="#bbb" stroke-dasharray="4"/>
<text x="{pad + 4}" y="{Hp - pad - total * (Hp - 2 * pad) / mx - 4:.1f}" font-size="10" fill="#888">総量 {total} PL</text>
<polyline points="{pline}" fill="none" stroke="#c7d8ee" stroke-width="3"/>
<polyline points="{aline}" fill="none" stroke="#2f7bd0" stroke-width="3"/>
{xlabels}
<text x="{Wp - pad - 150}" y="{pad}" font-size="11" fill="#8aa8cc">─ 計画(累積)</text>
<text x="{Wp - pad - 150}" y="{pad + 16}" font-size="11" fill="#2f7bd0">─ 実績(累積)</text>
</svg>
<p style="font-size:12px">実績が計画線より下にある日数分が遅れ。ロット別の内訳は 移管_マトリクス.xlsx を参照。</p>
</body></html>"""
    (out / "移管_バーンダウン.html").write_text(doc, encoding="utf-8")


# ---------- 統合スナップショット md ----------
def write_md(tasks, assigns, lots, quota, warns, today, out, name):
    wk_s = monday(today); wk_e = wk_s + dt.timedelta(days=6)
    late = [t for t in tasks if t["state"] != "完了" and t["end"] < today]
    this_wk = [t for t in tasks if t["state"] != "完了" and not (t["end"] < wk_s or t["start"] > wk_e)]
    ms = [t for t in tasks if t["ms"] and today <= t["end"] <= today + dt.timedelta(days=60)]
    done = sum(1 for t in tasks if t["state"] == "完了")

    def line(t, with_days=False):
        d = (f"(期限超過 {(today - t['end']).days}日)" if with_days
             else f"({t['start'].strftime('%m/%d')}〜{t['end'].strftime('%m/%d')})")
        return f"- {t['id']} {t['l3'] or t['l2']} / 担当: {t['own']} {d} 状態: {t['state']}" + \
               (f" — {t['note']}" if t["note"] else "")

    md = [f"# {name} スケジュール(最新スナップショット)",
          f"生成: {today} / 出典: タスク表・割当表・移管計画(このファイルは自動生成。直接編集しない)",
          f"進捗: タスク全{len(tasks)}件中 完了{done}件 / 遅延{len(late)}件", ""]
    if warns:
        md += ["## ⚠ 整合チェック(要修正)"] + [f"- {w}" for w in warns] + [""]
    md += ["## 遅延タスク(期限超過・未完了)"] + ([line(t, True) for t in late] or ["- なし"]) + [""]
    md += [f"## 今週のタスク({wk_s.strftime('%m/%d')}〜{wk_e.strftime('%m/%d')})"] + \
          ([line(t) for t in this_wk] or ["- なし"]) + [""]
    md += ["## 直近60日のマイルストーン"] + \
          ([f"- {t['end'].strftime('%m/%d')} ◆ {t['l3'] or t['l2']} ({t['id']})" for t in ms] or ["- なし"]) + [""]
    if assigns:
        wk_a = [a for a in assigns if wk_s <= a["date"] <= wk_e]
        md += [f"## 今週の支援者({wk_s.strftime('%m/%d')}〜{wk_e.strftime('%m/%d')})"]
        md += ([f"- {a['date'].strftime('%m/%d')} {a['who']}({a['org']}) → {a['place']} {a['work']}"
                + (f" [{a['task']}]" if a['task'] else "") for a in sorted(wk_a, key=lambda x: (x["date"], x["who"]))]
               or ["- なし"])
        counts = {}
        for a in assigns:
            counts[a["who"]] = counts.get(a["who"], 0) + 1
        md += ["", "### 支援者の割当日数(全期間)"] + \
              [f"- {w}: {n}日" for w, n in sorted(counts.items(), key=lambda x: -x[1])] + [""]
    if lots:
        total = sum(l["total"] for l in lots.values())
        act = sum(q["act"] or 0 for q in quota)
        plan_to_date = sum(q["plan"] for q in quota if q["date"] <= today)
        md += ["## 移管の進捗",
               f"- 総量 {total} PL / 実績 {act} PL ({(act / total * 100) if total else 0:.0f}%)",
               f"- 今日までの計画 {plan_to_date} PL に対し実績 {act} PL"
               f"({'遅れ ' + str(plan_to_date - act) + ' PL' if act < plan_to_date else '計画どおり'})"]
        for key, lot in sorted(lots.items()):
            a = sum(q["act"] or 0 for q in quota if q["key"] == key)
            md += [f"- {'/'.join(key)}: {a}/{lot['total']} PL"]
        md += [""]
    (out / "スケジュール_最新.md").write_text("\n".join(md), encoding="utf-8")


# ---------- 変更履歴(3表まとめてdiff) ----------
def write_changelog(tasks, assigns, quota, today, out):
    snap_p = out / ".snapshot.json"
    log_p = out / "変更履歴.md"
    cur = {
        "tasks": {t["id"]: {"l3": t["l3"], "start": str(t["start"]), "end": str(t["end"]),
                            "state": t["state"]} for t in tasks},
        "assigns": {f"{a['date']}|{a['who']}": f"{a['place']} {a['work']}".strip() for a in assigns},
        "quota": {f"{'/'.join(q['key'])}|{q['date']}": [q["plan"], q["act"]] for q in quota},
    }
    lines = []
    if snap_p.exists():
        prev = json.loads(snap_p.read_text(encoding="utf-8"))
        pt = prev.get("tasks", prev)  # 旧形式(タスクのみ)互換
        for i in sorted(set(cur["tasks"]) - set(pt)):
            c = cur["tasks"][i]; lines.append(f"- タスク追加: {i} {c['l3']} ({c['start']}〜{c['end']})")
        for i in sorted(set(pt) - set(cur["tasks"])):
            lines.append(f"- タスク削除: {i} {pt[i]['l3']}")
        for i in sorted(set(cur["tasks"]) & set(pt)):
            a, b = pt[i], cur["tasks"][i]
            if (a["start"], a["end"]) != (b["start"], b["end"]):
                lines.append(f"- 日程変更: {i} {b['l3']} {a['start']}〜{a['end']} → {b['start']}〜{b['end']}")
            if a["state"] != b["state"]:
                lines.append(f"- 状態変更: {i} {b['l3']} {a['state']} → {b['state']}")
        pa = prev.get("assigns", {})
        for k in sorted(set(cur["assigns"]) - set(pa)):
            lines.append(f"- 割当追加: {k.replace('|', ' ')} → {cur['assigns'][k]}")
        for k in sorted(set(pa) - set(cur["assigns"])):
            lines.append(f"- 割当取消: {k.replace('|', ' ')}")
        for k in sorted(set(cur["assigns"]) & set(pa)):
            if pa[k] != cur["assigns"][k]:
                lines.append(f"- 割当変更: {k.replace('|', ' ')} {pa[k]} → {cur['assigns'][k]}")
        pq = prev.get("quota", {})
        for k in sorted(set(cur["quota"]) & set(pq)):
            (p1, a1), (p2, a2) = pq[k], cur["quota"][k]
            if p1 != p2:
                lines.append(f"- 移管計画変更: {k.replace('|', ' ')} {p1} → {p2} PL")
            if a1 != a2:
                lines.append(f"- 移管実績記入: {k.replace('|', ' ')} {a2} PL")
        if lines:
            old = log_p.read_text(encoding="utf-8") if log_p.exists() else "# 変更履歴(自動生成)\n"
            log_p.write_text(old + f"\n## {today}\n" + "\n".join(lines) + "\n", encoding="utf-8")
            print(f"変更履歴: {len(lines)}件追記")
        else:
            print("変更履歴: 変更なし")
    else:
        log_p.write_text(f"# 変更履歴(自動生成)\n\n## {today}\n- 初回登録: タスク{len(cur['tasks'])}件"
                         f" / 割当{len(cur['assigns'])}件 / 移管計画{len(cur['quota'])}件\n", encoding="utf-8")
    snap_p.write_text(json.dumps(cur, ensure_ascii=False, indent=1), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src", help="tasks.csv 単体、またはプロジェクトフォルダ")
    ap.add_argument("--out", default=None)
    ap.add_argument("--today", default=None)
    a = ap.parse_args()
    src = pathlib.Path(a.src)
    today = parse_date(a.today) if a.today else dt.date.today()
    if src.is_dir():
        name = src.name
        tasks_p = src / "tasks.csv"
        if not tasks_p.exists():
            sys.exit(f"{src}/tasks.csv が見つからない")
        out = pathlib.Path(a.out) if a.out else src / "out"
        assigns = load_assignments(src / "assignments.csv") if (src / "assignments.csv").exists() else []
        lots, quota = (load_quota(src / "lots.csv", src / "quota.csv")
                       if (src / "lots.csv").exists() and (src / "quota.csv").exists() else ({}, []))
    else:
        name = src.stem.replace("_tasks", "").replace("tasks_", "")
        tasks_p = src
        out = pathlib.Path(a.out) if a.out else src.parent
        assigns, lots, quota = [], {}, []
    out.mkdir(parents=True, exist_ok=True)
    tasks = load_tasks(tasks_p)
    warns = cross_check(tasks, assigns, lots, quota)
    weeks = week_axis(tasks)
    write_xlsx(tasks, weeks, today, out, name)
    write_html(tasks, weeks, today, out, name)
    if assigns:
        write_assign_outputs(assigns, today, out)
    if lots:
        write_quota_outputs(lots, quota, today, out)
    write_md(tasks, assigns, lots, quota, warns, today, out, name)
    write_changelog(tasks, assigns, quota, today, out)
    for w in warns:
        print(f"⚠ {w}")
    print(f"OK: タスク{len(tasks)} / 割当{len(assigns)} / 移管計画{len(quota)} → {out}/")


if __name__ == "__main__":
    main()
