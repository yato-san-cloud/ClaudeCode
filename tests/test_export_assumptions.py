"""前提条件 as data, and the end of silent slide overflow.

Two defects blocked a real deliverable and were worked around instead of fixed:

  1. The 前提条件 of a proposal were a hardcoded constant, so the two 危険側
     (unsafe-side) assumptions an audit required could not be put on the 前提条件
     slide at all. They ended up spoken instead of written.
  2. PPTX text boxes had no autofit, so a slide silently ran off the page and the
     fix everyone reached for was rewriting the bullets shorter — i.e. the LAYOUT
     dictated the CONTENT.

These tests pin both fixes AND the promise that an export which supplies no new
data is unchanged. Fixtures are synthetic; no customer data is tracked here.
"""

from __future__ import annotations

import hashlib
import zipfile

import pytest

from whsim import export_doc
from whsim.export import textfit
from whsim.export._data import (
    INK,
    RED,
    _assumption_blocks,
    _cost_assumption_line,
    _normalize_assumptions,
)
from whsim.export.htmlviewer import build_viewer_html
from whsim.schema.model import WarehouseModel

# --- fixtures (synthetic) ----------------------------------------------------

KPIS = {
    "verdict": "現行体制で需要を充足できます（余力あり）。",
    "can_handle_demand": True,
    "throughput_per_hr": 208.6,
    "headcount": 200,
    "labour_rate_per_hr": 1_800,
    "capex_total": 24_000_000,
    "currency": "¥",
}

# The two 危険側 lines the audit required, plus one 注意. Verbatim-shaped, but
# about a made-up line so nothing customer-specific lands in the repo.
UNSAFE = [
    {"text": "本表は積み付けを拘束条件から除外している。荷降ろしバッファを置かない場合、"
             "ライン全体は208.6件/時まで低下する。", "level": "危険側"},
    {"text": "200名のピッカーは能力測定上の設定値であり、要員計画ではない。",
     "level": "危険側"},
    {"text": "保管坪単価は関東圏の相場を用いた概算である。", "level": "注意"},
]

LONG_INSIGHTS = [
    {"severity": "danger", "title": f"指摘{i}：工程の滞留がボトルネックになっている",
     "fact": "稼働率が高止まりし、待ち行列が時間とともに伸びていく状況が観測されています。" * 3,
     "action": "設備を1台増設し、投入を平準化してください。"}
    for i in range(25)
]


# --- pptx helpers ------------------------------------------------------------

def _slides(path):
    from pptx import Presentation
    return list(Presentation(str(path)).slides)


def _runs(path) -> list:
    """Every (text, pt, bold, rgb) run in the deck, in reading order."""
    out = []
    for slide in _slides(path):
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            for para in shp.text_frame.paragraphs:
                for r in para.runs:
                    rgb = None
                    try:
                        rgb = tuple(r.font.color.rgb)
                    except Exception:  # noqa: BLE001 — theme colour, not RGB
                        rgb = None
                    out.append((r.text, r.font.size.pt if r.font.size else None,
                                bool(r.font.bold), rgb))
    return out


def _all_text(path) -> str:
    return "".join(t for t, *_ in _runs(path))


def _footer_runs(path) -> list:
    """Runs in the bottom 0.5in strip of every content slide (the cover's own
    methodology line sits mid-slide at a different size and is not a footer)."""
    from pptx.util import Inches
    out = []
    for slide in _slides(path):
        for shp in slide.shapes:
            if not shp.has_text_frame or shp.top < Inches(6.9):
                continue
            for para in shp.text_frame.paragraphs:
                for r in para.runs:
                    out.append((r.text, r.font.size.pt if r.font.size else None))
    return out


def _room_in(slide, shp) -> float:
    """Vertical space a shape's text may really occupy: down to the top of the
    next shape below it (or the slide foot). The cover's title and 提案日 blocks
    deliberately sit in a gap larger than their own frame."""
    limits = [o.top for o in slide.shapes
              if o is not shp and o.top > shp.top
              and o.left < shp.left + shp.width and shp.left < o.left + o.width]
    limit = min(limits) if limits else int(7.5 * 914400)
    return max(shp.height, limit - shp.top) / 914400


def _overflowing_boxes(path) -> list[tuple[int, str, float, float]]:
    """Re-measure every rendered text box against its own frame.

    This closes the loop: the deck is read back from disk and measured with the
    same metrics the builder used, so "no overflow" is asserted about the file
    that would reach the customer, not about an intention."""
    bad = []
    for i, slide in enumerate(_slides(path)):
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            paras = []
            for para in shp.text_frame.paragraphs:
                if not para.runs:
                    continue
                pt = para.runs[0].font.size
                paras.append({"text": "".join(r.text for r in para.runs),
                              "pt": pt.pt if pt else 18.0})
            if not paras:
                continue
            w_in = shp.width / 914400
            # The cover's boxes are sized smaller than the gap they sit in; every
            # other box is measured against its own frame.
            h_in = _room_in(slide, shp) if i == 0 else shp.height / 914400
            used = textfit.block_height_in(paras, w_in)
            if used > h_in + 1e-9:
                bad.append((i, paras[0]["text"][:24], used, h_in))
    return bad


def _part_hashes(path) -> dict:
    """SHA-256 of every part inside the .pptx zip.

    Raw file bytes are NOT comparable: python-pptx stamps the current time into
    each zip entry, so two identical decks built a second apart differ. The parts
    themselves are the document."""
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in sorted(z.namelist())}


# --- textfit: the metric the whole contract rests on -------------------------

def test_textfit_full_width_vs_halfwidth():
    assert textfit.text_width_em("前提条件") == pytest.approx(4.0)
    assert textfit.text_width_em("abcd") < 4.0
    assert textfit.char_width_em("・") == 1.0  # CJK punctuation is full width


def test_textfit_wrapped_lines_counts_wraps():
    # 10 CJK chars in a box 5 em wide -> 2 lines.
    width_in = 5 * 12 / 72  # 5 em at 12pt
    assert textfit.wrapped_lines("あ" * 10, 12, width_in) == 2
    assert textfit.wrapped_lines("あ", 12, width_in) == 1
    assert textfit.wrapped_lines("あ\nい", 12, width_in) == 2  # honours newlines


def test_textfit_fit_scale_is_exactly_one_when_it_fits():
    paras = [{"text": "短い行", "pt": 16}]
    assert textfit.fit_scale(paras, 12.1, 5.0) == 1.0


def test_textfit_fit_scale_shrinks_and_never_below_floor():
    paras = [{"text": "長い行です。" * 200, "pt": 16}]
    scale = textfit.fit_scale(paras, 12.1, 2.0)
    assert textfit.MIN_SCALE <= scale < 1.0
    # ... and the result genuinely fits (or sits at the floor, to be paginated).
    assert (not textfit.overflows(paras, 12.1, 2.0, scale)
            or scale == pytest.approx(textfit.MIN_SCALE))


def test_textfit_paginate_keeps_every_paragraph():
    paras = [{"text": f"行{i}" + "あ" * 40, "pt": 16} for i in range(40)]
    pages = textfit.paginate(paras, 12.1, 5.0, textfit.MIN_SCALE)
    assert len(pages) > 1
    flat = [p for page in pages for p in page]
    assert flat == paras  # order preserved, nothing dropped
    for page in pages:
        assert not textfit.overflows(page, 12.1, 5.0, textfit.MIN_SCALE) or len(page) == 1


def test_textfit_paginate_keeps_headline_with_its_detail():
    paras = []
    for i in range(12):
        paras.append({"text": f"■ 指摘{i}" + "あ" * 30, "pt": 17, "keep_next": True})
        paras.append({"text": "　→ " + "い" * 60, "pt": 13})
    pages = textfit.paginate(paras, 12.1, 5.6, 1.0)
    for page in pages:
        assert not page[-1].get("keep_next"), "a headline was orphaned at a page break"


def test_textfit_fit_line_shrinks_then_elides():
    same, pt = textfit.fit_line("短い脚注", 12.3, 9.0)
    assert (same, pt) == ("短い脚注", 9.0)  # untouched -> byte-identical footers
    text, pt = textfit.fit_line("長い脚注。" * 40, 12.3, 9.0)
    assert pt < 9.0
    assert textfit.text_width_em(text) <= (12.3 * 72.0) / pt + 1e-9


# --- the assumptions API -----------------------------------------------------

def test_normalize_accepts_every_documented_shape():
    assert _normalize_assumptions(None) == ([], True)
    assert _normalize_assumptions("一行") == ([{"text": "一行", "level": "info"}], True)
    items, keep = _normalize_assumptions(["A", {"text": "B", "level": "危険側"},
                                          ("C", "注意")])
    assert keep is True
    assert [i["level"] for i in items] == ["info", "danger", "caution"]
    # envelope: replace drops the built-in cost line
    items, keep = _normalize_assumptions({"lines": ["X"], "replace": True})
    assert keep is False and items == [{"text": "X", "level": "info"}]
    # a single item dict is not mistaken for the envelope
    items, keep = _normalize_assumptions({"text": "Y", "level": "danger"})
    assert keep is True and items[0]["level"] == "danger"


def test_normalize_is_defensive():
    assert _normalize_assumptions(123) == ([], True)          # not iterable
    assert _normalize_assumptions([None, "", {}]) == ([], True)  # empties dropped
    # unknown levels degrade to 情報 rather than raising
    items, _ = _normalize_assumptions([{"text": "Z", "level": "なにか"}])
    assert items[0]["level"] == "info"


def test_default_blocks_are_the_historical_lines():
    """The default MUST still be the string every existing deck already prints."""
    blocks = _assumption_blocks(KPIS, "実データ 62%")
    assert [b["text"] for b in blocks] == [
        "本提案は離散事象シミュレーション(SimPy)に基づく概算です。",
        "実データ 62%",
        "前提条件: 人件費 ¥1,800/人時 ・ AGV投資 ¥24,000,000 ・ 36ヶ月償却",
    ]
    assert {b["level"] for b in blocks} == {"info"}
    assert {b["color"] for b in blocks} == {INK}


def test_extra_lines_are_appended_and_badged():
    blocks = _assumption_blocks(KPIS, "実データ 62%", UNSAFE)
    texts = [b["text"] for b in blocks]
    assert _cost_assumption_line(KPIS) in texts       # built-ins still there
    assert texts[3].startswith("【危険側】")
    assert texts[5].startswith("【注意】")
    assert blocks[3]["color"] == RED


def test_replace_drops_the_irrelevant_cost_line():
    blocks = _assumption_blocks(
        KPIS, "", {"lines": UNSAFE, "replace": True})
    texts = [b["text"] for b in blocks]
    assert not any(t.startswith("前提条件: 人件費") for t in texts)
    assert len(texts) == 1 + len(UNSAFE)  # methodology + the caller's lines


def test_model_settings_supply_the_assumptions():
    """When no argument is given the MODEL's own assumptions are used, so the
    deck states what the model assumes rather than what the template hardcoded."""
    class _S:
        assumptions = ["モデル由来の前提"]

    class _M:
        settings = _S()

    texts = [b["text"] for b in _assumption_blocks(KPIS, "", None, model=_M())]
    assert "モデル由来の前提" in texts
    # an explicit argument wins over the model
    texts = [b["text"] for b in _assumption_blocks(KPIS, "", ["引数の前提"], model=_M())]
    assert "引数の前提" in texts and "モデル由来の前提" not in texts


# --- PPTX: the assumptions reach the 前提条件 slide ---------------------------

def test_pptx_unsafe_assumptions_render_on_the_assumptions_slide(tmp_path):
    out = export_doc.build_pptx(KPIS, "テスト倉庫", "実データ 62%", None,
                                tmp_path / "a.pptx", assumptions=UNSAFE)
    runs = _runs(out)
    text = "".join(t for t, *_ in runs)
    assert "積み付けを拘束条件から除外" in text
    assert "能力測定上の設定値" in text
    # ... on the 前提条件 slide itself, not squeezed onto some other slide.
    last = _slides(out)[-1]
    slide_text = "".join(r.text for shp in last.shapes if shp.has_text_frame
                         for p in shp.text_frame.paragraphs for r in p.runs)
    assert "前提条件とデータ出所" in slide_text
    assert "【危険側】" in slide_text and "【注意】" in slide_text
    # and distinctly: red + bold, so it cannot be mistaken for boilerplate.
    danger = [r for r in runs if "【危険側】" in r[0]]
    assert danger and all(r[2] and r[3] == RED for r in danger)


def test_pdf_and_pptx_state_the_same_assumptions(tmp_path, monkeypatch):
    """A fact on the deck but not in the PDF is its own failure."""
    captured = {}
    from reportlab.platypus import SimpleDocTemplate
    orig = SimpleDocTemplate.build

    def _spy(self, story, *a, **kw):
        captured["story"] = list(story)
        return orig(self, story, *a, **kw)

    monkeypatch.setattr(SimpleDocTemplate, "build", _spy)
    export_doc.build_pdf(KPIS, "テスト倉庫", "実データ 62%", None,
                         tmp_path / "a.pdf", assumptions=UNSAFE)
    pdf_text = "".join(getattr(p, "text", "") for p in captured["story"])
    pptx = export_doc.build_pptx(KPIS, "テスト倉庫", "実データ 62%", None,
                                 tmp_path / "a.pptx", assumptions=UNSAFE)
    deck_text = _all_text(pptx)
    for blk in _assumption_blocks(KPIS, "実データ 62%", UNSAFE):
        assert blk["text"] in deck_text
        assert blk["text"] in pdf_text


def test_pdf_escapes_markup_in_caller_text(tmp_path):
    """reportlab parses its Paragraph text as markup — an ``&`` must not kill
    the export (a 前提 is written by a human, not by us)."""
    out = export_doc.build_pdf(
        KPIS, "テスト倉庫", "", None, tmp_path / "esc.pdf",
        assumptions=["A&B社の実績値を <参考> として用いた"])
    assert out.is_file() and out.stat().st_size > 3_000


def test_viewer_carries_the_same_assumptions():
    html = build_viewer_html(WarehouseModel(), kpis=KPIS,
                             provenance_summary="実データ 62%", assumptions=UNSAFE)
    assert "前提条件" in html
    assert "【危険側】" in html and "lv-danger" in html
    assert "積み付けを拘束条件から除外" in html
    # a viewer with no run keeps its previous shape (no 前提 built from 「—」)
    assert "前提条件" not in build_viewer_html(WarehouseModel())


# --- PPTX: overflow is impossible now ----------------------------------------

def test_pptx_long_content_never_overflows(tmp_path):
    """The defect: ~1,900 characters of required content ran off the slide, and
    the 'fix' was rewriting the bullets. Now it shrinks, then paginates."""
    long_assumptions = [{"text": f"前提{i}：" + "この試算は特定の条件下でのみ成立します。" * 4,
                         "level": "危険側" if i % 3 == 0 else ""} for i in range(14)]
    out = export_doc.build_pptx(KPIS, "長文テスト", "実データ 62%", None,
                                tmp_path / "long.pptx",
                                insights=LONG_INSIGHTS,
                                assumptions=long_assumptions)
    assert _overflowing_boxes(out) == []
    text = _all_text(out)
    # Nothing was dropped to make it fit — every supplied line is in the deck.
    for i in range(len(LONG_INSIGHTS)):
        assert f"指摘{i}：" in text
    for i in range(len(long_assumptions)):
        assert f"前提{i}：" in text
    # ... and the overflow went onto 「（続き）」 slides rather than off the page.
    assert "（続き）" in text


def test_pptx_overflow_regression_with_the_original_signature(tmp_path):
    """The narrowest statement of defect 2, using only the pre-existing API:
    a recommendation list that is too long for its box must not run off the
    slide. (This one fails against the previous revision.)"""
    out = export_doc.build_pptx(KPIS, "回帰", "実データ 62%", None,
                                tmp_path / "reg.pptx", insights=LONG_INSIGHTS)
    assert _overflowing_boxes(out) == []


def test_pptx_moderate_overflow_shrinks_instead_of_paginating(tmp_path):
    """Shrink-to-fit is tried FIRST: a slide that is only slightly too full stays
    one slide, at a smaller size, rather than splitting mid-argument."""
    insights = LONG_INSIGHTS[:9]
    out = export_doc.build_pptx(KPIS, "やや長い", "", None, tmp_path / "mid.pptx",
                                insights=insights)
    assert _overflowing_boxes(out) == []
    text = _all_text(out)
    assert "（続き）" not in text
    body = [r for r in _runs(out) if r[0].startswith("■ 指摘")]
    assert body and all(r[1] < 17 for r in body)  # authored at 17pt, shrunk


def test_pptx_long_cover_name_and_long_tile_value_shrink(tmp_path):
    """The other silent overflows on the same deck: a long 倉庫名 ran into the
    提案日 block, and a long ボトルネック name ran out of its tile."""
    kpis = dict(KPIS, bottleneck_jp="コンベア合流部（スパー3・シュート滞留）",
                verdict="現行体制では当日物量を充足できません。" * 3)
    out = export_doc.build_pptx(
        kpis, "関東広域物流センター 第2期 自動化検討（常温・冷蔵併設）" * 2,
        "実データ 62%", None, tmp_path / "cover.pptx")
    assert _overflowing_boxes(out) == []


def test_pptx_footer_is_immune_to_long_assumptions(tmp_path):
    """The footer is the SHORT form and must stay one line on the slide edge —
    the two uses of the methodology string are separate on purpose."""
    plain = export_doc.build_pptx(KPIS, "A", "実データ 62%", None,
                                  tmp_path / "p.pptx")
    loud = export_doc.build_pptx(KPIS, "A", "実データ 62%", None,
                                 tmp_path / "l.pptx",
                                 assumptions=[{"text": "長い前提。" * 80,
                                               "level": "危険側"}])
    foot = "本提案は離散事象シミュレーション(SimPy)に基づく ／ 実データ 62%"
    for path in (plain, loud):
        runs = _footer_runs(path)
        assert runs, "the short methodology footer is missing"
        assert all(r == (foot, 9.0) for r in runs)  # short form, unshrunk
    assert _footer_runs(plain) == _footer_runs(loud)
    assert _overflowing_boxes(loud) == []


def test_pptx_footer_shrinks_for_a_very_long_provenance(tmp_path):
    out = export_doc.build_pptx(KPIS, "A", "実データ " + "出所" * 120, None,
                                tmp_path / "prov.pptx")
    foots = _footer_runs(out)
    assert foots, "no footer run found"
    assert all(r[1] < 9.0 for r in foots)  # shrunk below the authored size
    # forced onto one line: shrunk, and elided only if that was not enough
    assert all(textfit.wrapped_lines(r[0], r[1], 12.3) == 1 for r in foots)


# --- byte-identity: an export that supplies nothing new is unchanged ---------

def test_supplying_no_assumptions_is_byte_identical(tmp_path):
    """Guards the promise proven offline against the previous revision: with no
    new data, every part of the package is identical."""
    a = export_doc.build_pptx(KPIS, "同一性", "実データ 62%", None,
                              tmp_path / "a.pptx", insights=LONG_INSIGHTS[:2])
    b = export_doc.build_pptx(KPIS, "同一性", "実データ 62%", None,
                              tmp_path / "b.pptx", insights=LONG_INSIGHTS[:2],
                              assumptions=None)
    assert _part_hashes(a) == _part_hashes(b)


def test_default_deck_keeps_its_authored_font_sizes(tmp_path):
    """A deck that fits is NOT rescaled — shrink-to-fit must be invisible until
    it is needed, or every existing export would silently change."""
    out = export_doc.build_pptx(KPIS, "既定", "実データ 62%", None,
                                tmp_path / "d.pptx", insights=LONG_INSIGHTS[:2])
    assumption_runs = [r for r in _runs(out) if r[0].startswith("・ 本提案は")]
    assert assumption_runs and assumption_runs[0][1] == 16.0
    rec_runs = [r for r in _runs(out) if r[0].startswith("■ ")]
    assert rec_runs and rec_runs[0][1] == 17.0
