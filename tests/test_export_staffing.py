"""Tests for the 人員配置と工程フロー section added to the PPTX + PDF proposal.

Covers: the shared derivation helper, the section rendering in both builders when
the model carries demand, its omission on a bare default model (no orders), and
``model=None`` back-compatibility of the builder signatures.
"""

from __future__ import annotations

import re

import pytest

from whsim import export_doc
from whsim.export._data import _staffing_section
from whsim.schema.model import Order, OrderLine, WarehouseModel, WorkProcess

CJK_RE = re.compile(r"[぀-ヿ一-鿿]")


# --- model builders ----------------------------------------------------------

def _demand_model(*, batch=False, custom=False) -> WarehouseModel:
    """A model with real outbound demand (spread across the working hours) so the
    staffing derivation yields positive volumes. Optionally sets a batch schedule
    and a custom work-process master."""
    m = WarehouseModel(meta={"name": "テスト物流センター"})
    orders = []
    for i in range(240):
        h = 9 + (i % 8)  # 9..16
        orders.append(Order(order_id=f"O{i:04d}", arrival_s=float(h * 3600),
                            lines=[OrderLine(sku=f"S{i % 25:03d}", qty=1 + i % 4)]))
    m.orders.outbound = orders
    if batch:
        m.settings.batch_schedule = {"出荷": [{"hour": 9, "pct": 70},
                                              {"hour": 13, "pct": 30}]}
    if custom:
        m.process.work_processes = [
            WorkProcess(id="荷受", section="入荷", driver="in_lines", prod=45,
                        unit="行/h", depends=[]),
            WorkProcess(id="ピッキング", section="出荷", driver="out_lines",
                        prod=55, unit="行/h", depends=["荷受"]),
            WorkProcess(id="出荷検品", section="出荷", driver="out_orders",
                        prod=90, unit="件/h", depends=["ピッキング"]),
        ]
    return m


# --- PPTX / text helpers -----------------------------------------------------

def _pptx_text(path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    parts: list[str] = []
    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.has_text_frame:
                for para in shp.text_frame.paragraphs:
                    parts += [r.text for r in para.runs]
            if shp.has_table:
                for row in shp.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            parts += [r.text for r in para.runs]
    return "".join(parts)


def _pptx_slides(path) -> int:
    from pptx import Presentation
    return len(Presentation(str(path)).slides)


def _pdf_pages(path) -> int:
    return len(re.findall(rb"/Type\s*/Page[^s]", path.read_bytes()))


# --- helper (pure derivation) ------------------------------------------------

def test_staffing_helper_with_demand_returns_full_section():
    staff = _staffing_section(_demand_model(batch=True))
    assert staff is not None
    labels = [t[0] for t in staff["tiles"]]
    assert labels == ["総工数", "ピーク人数", "終了時刻"]
    assert "判定" in staff["verdict"]
    assert isinstance(staff["verdict_ok"], bool)
    assert staff["flow_header"] == ["工程", "区分", "生産性", "依存"]
    assert staff["flow_rows"]  # at least one process row
    # every row has the 4 columns
    assert all(len(r) == 4 for r in staff["flow_rows"])
    # batch summary present because a schedule was set
    assert staff["batch_line"] and "バッチ投入" in staff["batch_line"]
    assert "70%" in staff["batch_line"]


def test_staffing_helper_no_batch_line_when_unset():
    staff = _staffing_section(_demand_model(batch=False))
    assert staff is not None
    assert staff["batch_line"] is None


def test_staffing_helper_custom_processes_surface():
    staff = _staffing_section(_demand_model(custom=True))
    assert staff is not None
    ids = {r[0] for r in staff["flow_rows"]}
    assert {"荷受", "ピッキング", "出荷検品"} <= ids
    # a dependency is rendered (not the dash placeholder) on a downstream row
    pick = next(r for r in staff["flow_rows"] if r[0] == "ピッキング")
    assert "荷受" in pick[3]


def test_staffing_helper_bare_default_is_none():
    # A bare default model carries no orders -> section skipped.
    assert _staffing_section(WarehouseModel()) is None


def test_staffing_helper_model_none_is_none():
    assert _staffing_section(None) is None


# --- PPTX --------------------------------------------------------------------

def test_pptx_staffing_section_present_with_demand(tmp_path):
    m = _demand_model(batch=True, custom=True)
    out = export_doc.build_pptx(
        {}, "テスト物流センター", "", None, tmp_path / "staff.pptx", model=m)
    assert out.is_file()
    txt = _pptx_text(out)
    assert "人員配置と工程フロー" in txt
    assert "総工数" in txt and "ピーク人数" in txt and "終了時刻" in txt
    assert "判定" in txt
    assert "バッチ投入" in txt
    # 工程フロー table header cells.
    for h in ("工程", "区分", "生産性", "依存"):
        assert h in txt
    assert CJK_RE.search(txt)


def test_pptx_staffing_skipped_without_demand(tmp_path):
    base = export_doc.build_pptx(
        {}, "空モデル", "", None, tmp_path / "nodemand.pptx",
        model=WarehouseModel())
    assert "人員配置と工程フロー" not in _pptx_text(base)


def test_pptx_staffing_adds_a_slide(tmp_path):
    without = export_doc.build_pptx(
        {}, "N", "", None, tmp_path / "wo.pptx")
    with_model = export_doc.build_pptx(
        {}, "N", "", None, tmp_path / "wm.pptx", model=_demand_model())
    assert _pptx_slides(with_model) == _pptx_slides(without) + 1


def test_pptx_model_none_back_compat(tmp_path):
    # Explicit model=None must behave like the old signature (no section).
    out = export_doc.build_pptx(
        {}, "N", "", None, tmp_path / "none.pptx", model=None)
    assert out.is_file()
    assert "人員配置と工程フロー" not in _pptx_text(out)


# --- PDF ---------------------------------------------------------------------

def test_pdf_staffing_section_present_with_demand(tmp_path):
    m = _demand_model(batch=True)
    with_model = export_doc.build_pdf(
        {}, "テスト物流センター", "", None, tmp_path / "staff.pdf", model=m)
    without = export_doc.build_pdf(
        {}, "テスト物流センター", "", None, tmp_path / "plain.pdf")
    assert with_model.is_file()
    assert _pdf_pages(with_model) >= 1
    # The extra section makes the document strictly larger than the plain one.
    assert with_model.stat().st_size > without.stat().st_size


def test_pdf_staffing_skipped_without_demand(tmp_path):
    # Bare model -> same output as no model at all (section omitted, no crash).
    bare = export_doc.build_pdf(
        {}, "空モデル", "", None, tmp_path / "bare.pdf", model=WarehouseModel())
    plain = export_doc.build_pdf(
        {}, "空モデル", "", None, tmp_path / "plain.pdf")
    assert bare.is_file()
    assert bare.stat().st_size == pytest.approx(plain.stat().st_size, rel=0.02)


def test_pdf_model_none_back_compat(tmp_path):
    out = export_doc.build_pdf(
        {}, "N", "", None, tmp_path / "none.pdf", model=None)
    assert out.is_file() and out.stat().st_size > 3_000
