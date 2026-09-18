"""Round-2 tests for the commercial-grade proposal export (PPTX + PDF) and the
proposal PNG. Verifies multi-section output, editable text, CJK rendering, and
graceful degradation on empty/None inputs and degenerate geometry."""

from __future__ import annotations

import re

import numpy as np
import pytest

from whsim import export_doc
from whsim.render.png2d import render as render_png
from whsim.schema.model import WarehouseModel


# --- fixtures ----------------------------------------------------------------

FULL_KPIS = {
    "verdict": "現行体制で需要を充足できます（余力あり）。",
    "can_handle_demand": True,
    "throughput_per_hr": 182.5,
    "throughput_p5": 160.0,
    "throughput_p95": 205.0,
    "orders_completed": 1460,
    "orders_arrived": 1500,
    "completion_rate": 0.973,
    "bottleneck_jp": "梱包工程",
    "bottleneck_utilization": 0.82,
    "picker_utilization": 0.74,
    "packer_utilization": 0.82,
    "agv_utilization": 0.55,
    "n_pickers": 6,
    "n_packers": 4,
    "n_agvs": 3,
    "cycle_p50_s": 540.0,
    "cycle_p95_s": 1120.0,
    "walk_per_order_m": 88.0,
    "headcount": 10,
    "total_cost_per_order": 312.4,
    "monthly_cost": 4_680_000,
    "monthly_opex": 4_200_000,
    "payback_months": 14.2,
    "labour_rate_per_hr": 1_800,
    "capex_total": 24_000_000,
    "currency": "¥",
}

SCENARIOS = {
    "baseline": {"name": "現行", "description": "現行オペレーション",
                 "kpis": FULL_KPIS},
    "alternatives": [{
        "name": "AGV導入", "description": "ピッキングをAGV化",
        "kpis": {**FULL_KPIS, "throughput_per_hr": 240.0, "headcount": 6,
                 "total_cost_per_order": 268.0, "monthly_cost": 3_900_000,
                 "payback_months": 11.5},
    }],
}

INSIGHTS = [
    {"severity": "danger", "title": "梱包工程がボトルネック",
     "fact": "稼働率 <span class=\"num\">82</span>%。",
     "action": "梱包台を1台増設で改善を検討。"},
    {"severity": "warn", "title": "AGV稼働率に余地",
     "fact": "55%。", "action": "搬送ルートの見直しを検討。"},
]

PROVENANCE = "本提案の 62% はお客様提供データに基づいています（残りは業界標準値）。"
MODEL_NAME = "サンプル物流センター"

CJK_RE = re.compile(r"[぀-ヿ一-鿿]")


@pytest.fixture
def layout_png(tmp_path):
    """A real proposal PNG produced by png2d from the default template model."""
    m = WarehouseModel(meta={"name": MODEL_NAME})
    heat = np.zeros((8, 8))
    heat[3, 3] = 5.0
    out = render_png(m, heat, FULL_KPIS, PROVENANCE, tmp_path / "layout.png")
    return out


# --- helpers -----------------------------------------------------------------

def _pptx_text_runs(path) -> list[str]:
    from pptx import Presentation
    prs = Presentation(str(path))
    texts: list[str] = []
    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.has_text_frame:
                for para in shp.text_frame.paragraphs:
                    texts += [r.text for r in para.runs]
            if shp.has_table:
                for row in shp.table.rows:
                    for cell in row.cells:
                        for para in cell.text_frame.paragraphs:
                            texts += [r.text for r in para.runs]
    return texts


def _pptx_slide_count(path) -> int:
    from pptx import Presentation
    return len(Presentation(str(path)).slides)


def _pdf_page_count(path) -> int:
    data = path.read_bytes()
    # Count page objects without an external PDF library.
    return len(re.findall(rb"/Type\s*/Page[^s]", data))


# --- PPTX tests --------------------------------------------------------------

def test_pptx_full_data_is_editable_multi_section(tmp_path, layout_png):
    out = export_doc.build_pptx(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
        tmp_path / "full.pptx", scenarios=SCENARIOS, insights=INSIGHTS,
    )
    assert out.is_file()
    assert out.stat().st_size > 20_000  # non-trivial deck

    # Multiple slides incl. the optional scenario + recommendation sections.
    assert _pptx_slide_count(out) >= 6

    runs = _pptx_text_runs(out)
    assert len(runs) > 20  # native editable text, not one flat image
    joined = "".join(runs)
    # Cover, exec-summary headline KPIs, sections, recommendation, methodology.
    assert "提案書" in joined
    assert "エグゼクティブサマリー" in joined
    assert "処理能力" in joined
    assert "シナリオ比較（現行 vs 代替案）" in joined
    assert "ご提案" in joined
    assert "SimPy" in joined
    # Insight HTML span must be stripped to plain text.
    assert "<span" not in joined
    assert CJK_RE.search(joined)


def test_pptx_empty_kpis_does_not_raise(tmp_path):
    out = export_doc.build_pptx(
        None, "", "", None, tmp_path / "empty.pptx",
    )
    assert out.is_file()
    assert out.stat().st_size > 10_000
    # Cover + exec + layout + KPI + methodology still render (no scenarios).
    assert _pptx_slide_count(out) >= 5
    runs = _pptx_text_runs(out)
    assert any(r.strip() for r in runs)
    joined = "".join(runs)
    assert "—" in joined  # missing values rendered as the dash placeholder
    assert CJK_RE.search(joined)


def test_pptx_old_signature_still_works(tmp_path, layout_png):
    """The pre-existing positional call (no new kwargs) must keep working."""
    out = export_doc.build_pptx(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png, tmp_path / "compat.pptx",
    )
    assert out.is_file() and out.stat().st_size > 15_000


# --- PDF tests ---------------------------------------------------------------

def test_pdf_full_data_multi_page_with_cjk(tmp_path, layout_png):
    out = export_doc.build_pdf(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
        tmp_path / "full.pdf", scenarios=SCENARIOS, insights=INSIGHTS,
    )
    assert out.is_file()
    assert out.stat().st_size > 8_000
    assert _pdf_page_count(out) > 1


def test_pdf_empty_kpis_does_not_raise(tmp_path):
    out = export_doc.build_pdf(
        None, "", None, None, tmp_path / "empty.pdf",
    )
    assert out.is_file()
    assert out.stat().st_size > 3_000
    assert _pdf_page_count(out) >= 1


def test_pdf_old_signature_still_works(tmp_path, layout_png):
    out = export_doc.build_pdf(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png, tmp_path / "compat.pdf",
    )
    assert out.is_file() and out.stat().st_size > 5_000


# --- scenarios passed as a plain list ----------------------------------------

def test_scenarios_accepts_plain_list(tmp_path, layout_png):
    scen_list = [
        {"name": "現行", "kpis": FULL_KPIS},
        {"name": "案A", "kpis": {**FULL_KPIS, "throughput_per_hr": 300.0}},
    ]
    out = export_doc.build_pptx(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
        tmp_path / "list.pptx", scenarios=scen_list,
    )
    assert out.is_file()
    assert "シナリオ比較" in "".join(_pptx_text_runs(out))


# --- 提案書ブランドテーマ (brand theme) -------------------------------------

def _png_logo(path, color=(0xE2, 0x23, 0x1A)):
    """A tiny real PNG to embed as a brand logo."""
    from PIL import Image
    Image.new("RGBA", (200, 60), (*color, 255)).save(path)
    return path


def test_brand_default_is_a_noop(tmp_path, layout_png):
    """An absent/empty brand must export exactly the (colour-identical) baseline."""
    plain = export_doc.build_pptx(FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
                                  tmp_path / "plain.pptx")
    default_brand = export_doc.build_pptx(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png, tmp_path / "brand0.pptx",
        brand={"accent_color": "#2383E2"})  # the built-in accent
    assert plain.is_file() and default_brand.is_file()
    # Same text content (no 御中 / 提案元 added), and both are valid decks.
    j = "".join(_pptx_text_runs(default_brand))
    assert "御中" not in j and "提案元" not in j
    assert _pptx_slide_count(default_brand) == _pptx_slide_count(plain)


def test_brand_name_and_color_pptx(tmp_path, layout_png):
    brand = {"company_name": "提案元ロジ", "client_name": "アクメ物流",
             "accent_color": "#E2231A", "footer_note": "担当 山田 / 03-0000"}
    out = export_doc.build_pptx(FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
                                tmp_path / "brand.pptx", brand=brand)
    assert out.is_file() and out.stat().st_size > 20_000
    j = "".join(_pptx_text_runs(out))
    assert "アクメ物流 御中" in j        # 宛先
    assert "提案元：提案元ロジ" in j       # 提案元
    assert "担当 山田 / 03-0000" in j      # footer note


def test_brand_name_and_color_pdf(tmp_path, layout_png):
    brand = {"company_name": "提案元ロジ", "client_name": "アクメ物流",
             "accent_color": "#E2231A"}
    out = export_doc.build_pdf(FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
                               tmp_path / "brand.pdf", brand=brand)
    assert out.is_file() and out.stat().st_size > 8_000


def test_brand_logo_embeds_on_cover_pptx(tmp_path, layout_png):
    logo = _png_logo(tmp_path / "logo.png")
    out = export_doc.build_pptx(FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
                                tmp_path / "logo.pptx",
                                brand={"logo_path": str(logo)})
    from pptx import Presentation
    cover = Presentation(str(out)).slides[0]
    pics = [sh for sh in cover.shapes if sh.shape_type == 13]  # 13 = PICTURE
    assert len(pics) == 1  # the brand logo landed on the cover


def test_brand_bogus_logo_does_not_crash(tmp_path, layout_png):
    """A missing path and a corrupt image must both degrade silently."""
    missing = export_doc.build_pptx(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png, tmp_path / "miss.pptx",
        brand={"logo_path": "/no/such/logo.png", "accent_color": "not-a-color"})
    assert missing.is_file()
    # a real file whose *content* is not an image
    corrupt = tmp_path / "corrupt.png"
    corrupt.write_bytes(b"this is not a png")
    out_pptx = export_doc.build_pptx(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png, tmp_path / "corrupt.pptx",
        brand={"logo_path": str(corrupt)})
    out_pdf = export_doc.build_pdf(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png, tmp_path / "corrupt.pdf",
        brand={"logo_path": str(corrupt)})
    assert out_pptx.is_file() and out_pdf.is_file()
    from pptx import Presentation
    pics = [sh for sh in Presentation(str(out_pptx)).slides[0].shapes
            if sh.shape_type == 13]
    assert len(pics) == 0  # corrupt image skipped, cover still renders


def test_brand_from_model_settings(tmp_path, layout_png):
    """When no explicit brand is passed, model.settings.brand is honoured."""
    m = WarehouseModel(meta={"name": MODEL_NAME})
    m.settings.brand.client_name = "モデル客先"
    m.settings.brand.company_name = "モデル提案元"
    out = export_doc.build_pptx(FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png,
                                tmp_path / "model_brand.pptx", model=m)
    j = "".join(_pptx_text_runs(out))
    assert "モデル客先 御中" in j and "提案元：モデル提案元" in j


def test_brand_xml_special_chars_pdf(tmp_path, layout_png):
    """Brand text with reportlab-markup metacharacters must not break the PDF."""
    out = export_doc.build_pdf(
        FULL_KPIS, MODEL_NAME, PROVENANCE, layout_png, tmp_path / "xml.pdf",
        brand={"company_name": "A & B <Corp>", "client_name": 'C<>&"D'})
    assert out.is_file() and out.stat().st_size > 5_000


# --- png2d tests -------------------------------------------------------------

def test_png2d_renders_legible_normal_geometry(tmp_path):
    m = WarehouseModel(meta={"name": MODEL_NAME})
    heat = np.zeros((10, 10))
    heat[5, 5] = 9.0
    out = render_png(m, heat, FULL_KPIS, PROVENANCE, tmp_path / "normal.png")
    assert out.is_file()
    assert out.stat().st_size > 10_000  # a real, non-empty image


def test_png2d_handles_degenerate_geometry(tmp_path):
    """Empty model (zero/None bounds) and empty kpis must not raise."""
    m = WarehouseModel()
    m.layout.bounds.width = 0.0
    m.layout.bounds.depth = 0.0
    out = render_png(m, np.zeros((4, 4)), {}, "", tmp_path / "degenerate.png")
    assert out.is_file()
    assert out.stat().st_size > 5_000
