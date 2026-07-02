import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from jancode_dimensions.cli import main  # noqa: E402
from jancode_dimensions.excel_io import ExcelSheet  # noqa: E402


def _make_xlsx(path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_find_jan_column_auto(tmp_path):
    path = str(tmp_path / "in.xlsx")
    _make_xlsx(path, ["商品名", "JANコード", "在庫数"], [["A", 4901234567894, 1]])
    sheet = ExcelSheet(path)
    assert sheet.find_jan_column() == 2


def test_find_jan_column_missing_raises(tmp_path):
    path = str(tmp_path / "in.xlsx")
    _make_xlsx(path, ["コード", "数量"], [[123, 1]])
    sheet = ExcelSheet(path)
    with pytest.raises(ValueError):
        sheet.find_jan_column()


def test_ensure_output_columns_idempotent(tmp_path):
    path = str(tmp_path / "in.xlsx")
    _make_xlsx(path, ["JAN", "商品名"], [[4901234567894, "A"]])
    sheet = ExcelSheet(path)
    first = sheet.ensure_output_columns()
    second = sheet.ensure_output_columns()
    assert first == second  # 2回呼んでも同じ列に割り当てられ、列は増えない


def test_cli_end_to_end(tmp_path):
    master = tmp_path / "master.csv"
    master.write_text(
        "JANコード,商品名,幅,奥行,高さ\n4901234567894,マスタ商品A,30,20,10\n",
        encoding="utf-8",
    )
    xlsx = str(tmp_path / "in.xlsx")
    _make_xlsx(
        xlsx,
        ["JANコード", "商品名", "在庫数"],
        [
            [4901234567894, "うちの商品A", 10],
            [4999999999999, "うちの商品B", 2],
        ],
    )
    out = str(tmp_path / "out.xlsx")
    cache = str(tmp_path / "cache.sqlite")

    rc = main([
        "process", xlsx,
        "--master", str(master),
        "--output", out,
        "--cache", cache,
    ])
    assert rc == 0

    ws = openpyxl.load_workbook(out).active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])

    # 元の商品名列は保持され、照合商品名は別列に入る
    i_name = header.index("商品名")
    i_match = header.index("照合商品名")
    assert rows[1][i_name] == "うちの商品A"
    assert rows[1][i_match] == "マスタ商品A"

    # サイズと三辺合計が転記される
    assert float(rows[1][header.index("三辺合計(cm)")]) == 60.0
    assert rows[1][header.index("取得元")] == "local"

    # マスタに無いJANは未取得として記録される
    assert rows[2][i_name] == "うちの商品B"
    assert rows[2][header.index("備考")] == "サイズ取得できず"

    # 再実行: 列が増えず、キャッシュから解決される
    rc = main([
        "process", out,
        "--master", str(master),
        "--output", out,
        "--cache", cache,
    ])
    assert rc == 0
    ws2 = openpyxl.load_workbook(out).active
    assert ws2.max_column == len(header)
    rows2 = list(ws2.iter_rows(values_only=True))
    assert rows2[1][header.index("備考")] == "キャッシュ"
